"""Create a self-contained standard/optimized 59:50 timer Excel comparison."""
from collections import Counter
import importlib.util
import json
import math
from pathlib import Path
import statistics

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE=Path(__file__).resolve().parent
TIMER=HERE.parent
original=json.loads((TIMER/'witness.json').read_text())
run=json.loads((TIMER/'optimization/optimized.json').read_text())['runs'][1]
setup=json.loads((TIMER/'optimization/setups.json').read_text())['runs'][1]
spec=importlib.util.spec_from_file_location('timer_physics',TIMER/'verify.py')
P=importlib.util.module_from_spec(spec);spec.loader.exec_module(P)
label=lambda t:f'{t//6:02d}:{t%6*10:02d}'
htm=lambda w:len(w.split())
qtm=lambda w:sum(2 if t.endswith('2') else 1 for t in w.split())
raw=[w['word'] for w in original['witnesses']]
standard=[raw[0]]+[P.inverse(raw[t-1])+' '+raw[t] for t in range(1,360)]
scenarios=[('Standard',original['printing'],standard,raw[-1]),
           ('Direct baseline',original['printing'],run['baselineDirect']['words'],None),
           ('Optimized',run['printing'],run['optimized']['words'],setup['maximumWord'])]
# Revalidate every exported route and its inverse, using the independent engine.
for name,printing,words,maximum in scenarios:
    source=tuple((s['id'],s['cubie'],tuple(s['p']),tuple(s['n']),tuple(s['up']),s['image']) for s in printing)
    state=source;states=[]
    for tick,word in enumerate(words):
        state=P.apply(state,word);P.show(state,tick//6,tick%6*10);states.append(state)
    if maximum is not None:assert P.apply(source,maximum)==states[-1]
    for tick in range(359,0,-1):
        state=P.apply(state,P.inverse(words[tick]));assert state==states[tick-1]
    assert P.apply(state,P.inverse(words[0]))==source

wb=Workbook();overview=wb.active;overview.title='Overview'
header_fill=PatternFill('solid',fgColor='17365D')
role_colors=['D9EAF7','D9EAD3','FFF2CC','FCE4D6']

def table(title,headers,rows,widths=None):
    ws=wb.create_sheet(title);ws.append(headers)
    for row in rows:ws.append(row)
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]:cell.fill=header_fill;cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(wrap_text=True)
    for col,width in enumerate(widths or [18]*len(headers),1):ws.column_dimensions[get_column_letter(col)].width=width
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    ws.print_title_rows='1:1'
    return ws

summary=[]
for name,printing,words,maximum in scenarios:
    lengths=[htm(w) for w in words[1:]];quarters=[qtm(w) for w in words[1:]]
    summary.append([name,360,359,min(lengths),statistics.median(lengths),sorted(lengths)[math.ceil(.95*359)-1],max(lengths),statistics.mean(lengths),sum(lengths),statistics.mean(quarters),max(quarters),sum(quarters),htm(words[0]),htm(maximum) if maximum else 'Not exported'])
overview.append(['59:50 timer — standard and optimized fixed prints'])
overview.append(['Range','00:00–59:50; MM above SS; zero-padded; 10-second steps'])
overview.append(['Standard','Original fixed print and raw inverse-source/forward-target paths; no cancellation.'])
overview.append(['Direct baseline','Same standard print with direct paths, initial tie policy; isolates the routing baseline.'])
overview.append(['Optimized','Selected fixed print and canonical reversible routes.'])
overview.append(['Metric','HTM: quarter, inverse or half turn = 1. QTM: half turn = 2. Setup excluded from edge statistics.'])
overview.append(['Proof','All 360 displays and all 359 forward + 359 reverse edges per route physically replayed during export.'])
overview.append(['Limits','Global print optimality UNKNOWN. No speed/physical-fit guarantee. No wraparound.'])
overview.append(['Use','Start in an exact exported setup state. Reverse words restore canonical hidden state; arbitrary matching displays may differ.'])
overview.append(['Standard comparison','Raw hub paths are an upper-bound reference, not shortest paths. Fair direct baseline is included separately.'])
overview.append(['Prints','Assembly maps, not cut-out artwork. Digits are shown upright for legibility; arrows give print-up direction on each outward source face.'])
overview.append(['Sources','../witness.json; ../optimization/optimized.json; ../optimization/setups.json; ../optimization/REPORT.md'])
overview.append([])
headers=['Route','Displays','Edges','Min HTM','Median HTM','P95 HTM','Max HTM','Mean HTM','Total HTM','Mean QTM','Max QTM','Total QTM','00 setup HTM','59:50 setup HTM']
overview.append(headers)
for row in summary:overview.append(row)
overview.append([])
overview.append(['Optimized total reduction versus standard',f'=1-I17/I15'])
overview.append(['Optimized total reduction versus direct baseline',f'=1-I17/I16'])
for r in [19,20]:overview.cell(r,2).number_format='0.00%'
for r in range(15,18):
    for c in (8,10):overview.cell(r,c).number_format='0.000'
for cell in overview[14]:cell.fill=header_fill;cell.font=Font(color='FFFFFF',bold=True)
overview.column_dimensions['A'].width=48
for col in range(2,15):overview.column_dimensions[get_column_letter(col)].width=18
for r in range(2,13):
    overview.merge_cells(start_row=r,start_column=2,end_row=r,end_column=14)
    overview.cell(r,2).alignment=Alignment(wrap_text=True,vertical='center');overview.row_dimensions[r].height=32
for r in (19,20):overview.cell(r,1).alignment=Alignment(wrap_text=True);overview.row_dimensions[r].height=32
chart=BarChart();chart.title='Mean HTM per 10-second edge';chart.y_axis.title='HTM';chart.add_data(Reference(overview,min_col=8,min_row=14,max_row=17),titles_from_data=True);chart.set_categories(Reference(overview,min_col=1,min_row=15,max_row=17));overview.add_chart(chart,'A23')
overview.freeze_panes='B15'

for name,printing,words,maximum in scenarios:
    rows=[]
    for tick in range(1,360):
        rows.append([tick,label(tick-1),label(tick),words[tick],P.inverse(words[tick]),htm(words[tick]),qtm(words[tick]),'Decade rollover' if tick%60==0 else 'Minute rollover' if tick%6==0 else 'Seconds only'])
    table(name+' transitions',['Edge','From','To','Increment +10 sec','Decrement -10 sec','HTM','QTM','Category'],rows,[10,12,12,95,95,10,10,22])
comparison=[]
for tick in range(1,360):
    comparison.append([tick,label(tick-1),label(tick),htm(standard[tick]),htm(run['baselineDirect']['words'][tick]),htm(run['optimized']['words'][tick]),f'=D{tick+1}-F{tick+1}',f'=E{tick+1}-F{tick+1}'])
table('Edge comparison',['Edge','From','To','Standard HTM','Direct baseline HTM','Optimized HTM','Saved vs standard','Saved vs direct'],comparison)

faces={'U':((0,1,0),(1,0,0),(0,0,-1),(1,0)), 'L':((-1,0,0),(0,0,1),(0,1,0),(0,1)),
       'F':((0,0,1),(1,0,0),(0,1,0),(1,1)), 'R':((1,0,0),(0,0,-1),(0,1,0),(2,1)),
       'B':((0,0,-1),(-1,0,0),(0,1,0),(3,1)), 'D':((0,-1,0),(1,0,0),(0,0,1),(1,2))}
roles=['Minute tens','Minute units','Second tens','Second units']
dot=lambda a,b:sum(x*y for x,y in zip(a,b))
for name,printing in [('Standard',original['printing']),('Optimized',run['printing'])]:
    rows=[];net=wb.create_sheet(name+' print net')
    net.cell(1,1,f'{name} source-state print net — outward views')
    net.cell(2,1,'ASSEMBLY MAP: upright digit labels / print-up arrows / IDs. Not cut-out artwork; use arrows and sticker frames.')
    for c in range(1,14):net.column_dimensions[get_column_letter(c)].width=15
    for r in range(4,15):net.row_dimensions[r].height=65
    for face,(normal,right,up,(gx,gy)) in faces.items():
        start_col=1+gx*3;start_row=4+gy*4
        net.cell(start_row,start_col,face).font=Font(bold=True)
        net.row_dimensions[start_row].height=20
        for s in printing:
            if tuple(s['n'])!=normal:continue
            col=(dot(s['p'],right)+1)//2;row=(1-dot(s['p'],up))//2
            angle=round(math.degrees(math.atan2(dot(s['up'],right),dot(s['up'],up))))%360
            arrow={0:'↑',90:'→',180:'↓',270:'←'}[angle]
            cell=net.cell(start_row+1+row,start_col+col,f"{s['image']} {arrow}\nID {s['id']}")
            cell.font=Font(size=22,bold=True);cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True);cell.fill=PatternFill('solid',fgColor=role_colors[s['slot']])
            rows.append([s['id'],s['cubie'],face,row+1,col+1,roles[s['slot']],s['image'],angle,arrow,json.dumps(s['p']),json.dumps(s['n']),json.dumps(s['up'])])
    table(name+' stickers',['Sticker ID','Cubie','Face','Row','Column','Role','Digit','Clockwise degrees','Up arrow','Position','Normal','Up vector'],sorted(rows),[12,10,10,10,10,20,10,20,12,20,20,20])
    net.sheet_properties.pageSetUpPr.fitToPage=True;net.page_setup.orientation='landscape';net.page_setup.paperSize=net.PAPERSIZE_A4;net.page_setup.fitToWidth=1;net.page_setup.fitToHeight=1
    net.print_options.horizontalCentered=True;net.print_area='A1:L15'

setup_rows=[]
for name,printing,words,maximum in scenarios:
    setup_rows.append([name,'00:00',words[0],htm(words[0]),'Exact canonical zero from source-state print'])
    if maximum:setup_rows.append([name,'59:50',maximum,htm(maximum),'Exact canonical maximum from source-state print; not claimed shortest'])
table('Initialization',['Route','Display','Setup word','HTM','Note'],setup_rows,[22,14,100,10,85])
category_rows=[];hist=[]
for name,printing,words,maximum in scenarios:
    for kind in ['Seconds only','Minute rollover','Decade rollover']:
        lengths=[htm(words[t]) for t in range(1,360) if ('Decade rollover' if t%60==0 else 'Minute rollover' if t%6==0 else 'Seconds only')==kind]
        category_rows.append([name,kind,len(lengths),min(lengths),statistics.mean(lengths),max(lengths),sum(lengths)])
    hist.append(dict(Counter(htm(w) for w in words[1:])))
table('Rollover statistics',['Route','Category','Edges','Min HTM','Mean HTM','Max HTM','Total HTM'],category_rows)
table('Move histogram',['HTM']+[s[0] for s in scenarios],[[length]+[counts.get(length,0) for counts in hist] for length in range(1,22)])
for ws in wb:
    ws.sheet_view.showGridLines=False
    ws.sheet_properties.outlinePr.summaryRight=False
wb.calculation.fullCalcOnLoad=True
wb.calculation.forceFullCalc=True
wb.active=0
output=HERE/'timer-5950-standard-optimized.xlsx'
wb.save(output)
# Reopen and verify the actual exported binary, not just the in-memory workbook.
loaded=load_workbook(output,data_only=False)
assert len(loaded.sheetnames)==12
assert [loaded['Overview'].cell(r,9).value for r in (15,16,17)]==[6535,1709,1324]
for name,_,words,_ in scenarios:
    sheet=loaded[name+' transitions'];assert sheet.max_row==360
    for tick in range(1,360):
        assert sheet.cell(tick+1,4).value==words[tick]
        assert sheet.cell(tick+1,5).value==P.inverse(words[tick])
        assert sheet.cell(tick+1,6).value==htm(words[tick])
for name in ['Standard','Optimized']:assert loaded[name+' stickers'].max_row==25
assert loaded['Overview']['B19'].value=='=1-I17/I15'
assert loaded['Overview']['B20'].value=='=1-I17/I16'
print(f'Created and reopened {output}; 12 sheets; all three routes independently replayed in both directions.')
